"""Build MCB_HR_Flow_Guides.docx — the click-by-click guide.

For every step: the real screen with the button to click ringed in red, then
WHO is doing it, WHAT they click, and WHAT HAPPENS NEXT.
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BUILD = Path("/Data/odoo19_enterprise/custom-addons/mcb_hr/_build")
FLOWS = BUILD / "flows"
OUT = Path("/Data/odoo19_enterprise/mukticox/MCB_HR_Flow_Guides.docx")

# overridable so the same builder can produce the Accounts guide
HEADER_TEXT = "MCB \u00b7 HR — Click-by-Click Flow Guides"
DOC_TITLE = "HR — Click-by-Click Flow Guides"
DOC_SUB = ("Where to click, what opens next, and who does it —\n"
           "from the person who starts it to the person who approves it.")
INTRO_COUNT = "seven"

NAVY = RGBColor(0x1F, 0x36, 0x4D)
GREY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xC0, 0x28, 0x28)
GREEN = RGBColor(0x1E, 0x7A, 0x53)

# ------------------------------------------------------------------ flows ---
# Listed in the order the work actually happens at MCB.
FLOW_DEFS = [
    dict(
        key="requisition", title="Flow 1 — Asking for a New Staff Member",
        summary="A Project Manager needs someone new. He raises a Staff Requisition, "
                "HR checks it, the Chief Executive approves it, and the job circular "
                "is printed.",
        who=[("Project Manager", "Raises the requisition and confirms budget", "Part A"),
             ("HR", "Reviews it", "Part A"),
             ("Chief Executive", "Gives the final approval", "Part B"),
             ("HR", "Publishes it and prints the circular", "Part B")],
        logins=[("Project Manager / HR / CE", "admin", "admin")],
        parts={1: "PART A — Log in and raise the request", 8: "PART B — Approvals and publishing"},
        outcome="Approved, published requisition + a printed Job Circular",
        leads="Flow 2 — applications now arrive against this post.",
    ),
    dict(
        key="recruitment", title="Flow 2 — Applicants, Exams and the Merit List",
        summary="Candidates apply, are shortlisted, sit the written / computer / oral "
                "exams, and the system builds the merit list from the marks.",
        who=[("Recruitment Committee", "Shortlists and scores candidates", "Part A"),
             ("HR", "Issues admit cards", "Part A"),
             ("System", "Computes the totals and the ranking", "Part B")],
        logins=[("Committee / HR", "admin", "admin")],
        parts={1: "PART A — Log in, candidates and shortlisting", 7: "PART B — Exams, marks and ranking"},
        outcome="A ranked merit list with the selected candidate marked",
        leads="Flow 3 — the winning candidate becomes an employee.",
    ),
    dict(
        key="onboarding", title="Flow 3 — Hiring and Onboarding",
        summary="HR turns the selected candidate into an employee, completes the MCB "
                "profile, attaches the grade and launches the orientation plan.",
        who=[("HR", "Creates the employee from the applicant", "Part A"),
             ("HR", "Completes the MCB profile and contract", "Part B"),
             ("System", "Generates the employee code and the orientation tasks", "Part B")],
        logins=[("HR", "admin", "admin")],
        parts={1: "PART A — Log in, candidate to employee", 5: "PART B — Profile, grade and orientation"},
        outcome="An active employee with a code, a grade and an orientation plan running",
        leads="Flows 4–6 — he can now take leave, claim travel money and be paid.",
    ),
    dict(
        key="leave", title="Flow 4 — Applying for Leave and Getting it Approved",
        summary="An employee asks for 3 days off. His line manager approves it. "
                "The balance updates by itself.",
        who=[("Employee (Rashedul Karim)", "Applies for the leave", "Part A"),
             ("Line Manager", "Finds it waiting and approves it", "Part B"),
             ("System", "Drops the balance and records who approved", "Part B")],
        logins=[("Employee", "rashedul", "rashedul"),
                ("Manager / approver", "admin", "admin")],
        parts={1: "PART A — The employee logs in and applies",
               9: "PART B — The approver logs in and approves"},
        outcome="Leave approved · balance reduced automatically · full history kept",
        leads="Flow 6 — approved leave is taken into account by payroll.",
    ),
    dict(
        key="expense", title="Flow 5 — Claiming Travel Money (TA / DA)",
        summary="An employee returns from a field visit and claims his travel and food "
                "allowance. The rate comes from his grade, not from an argument.",
        who=[("Employee (Rashedul Karim)", "Records the trip and applies the per-diem", "Part A"),
             ("Head of Department", "Approves the claim", "Part B"),
             ("Finance", "Pays it, coded to the right donor", "after Part B")],
        logins=[("Employee", "rashedul", "rashedul"),
                ("Head of Department", "admin", "admin")],
        parts={1: "PART A — The employee logs in and claims",
               10: "PART B — The approver logs in and approves"},
        outcome="Approved claim, correct rate, charged to the right project and donor",
        leads="Finance reimburses and the cost lands in the donor report.",
    ),
    dict(
        key="payroll", title="Flow 6 — Running the Monthly Payroll",
        summary="Finance creates the month's batch, generates payslips, checks the "
                "calculation and produces the bank transfer file.",
        who=[("Finance", "Creates the batch and generates payslips", "Part A"),
             ("Finance", "Checks each payslip and confirms", "Part B"),
             ("Finance", "Generates the bank file", "Part B")],
        logins=[("Finance / Payroll Officer", "admin", "admin")],
        parts={1: "PART A — Log in and create the batch", 6: "PART B — Check, confirm and pay"},
        outcome="Confirmed payslips + one bank file that pays everybody",
        leads="Flow 7 — salary figures feed the final settlement when someone leaves.",
    ),
    dict(
        key="separation", title="Flow 7 — Resignation and Final Settlement",
        summary="A permanent staff member resigns with short notice. The system works "
                "out the shortfall, the settlement and the certificate.",
        who=[("Employee", "Submits the resignation", "Part A"),
             ("System", "Checks the notice period and computes salary in lieu", "Part A"),
             ("Manager → HR → CE", "Approve it", "Part B"),
             ("HR", "Settles, prints the certificate and closes access", "Part B")],
        logins=[("HR / Manager / CE", "admin", "admin")],
        parts={1: "PART A — Log in, resignation and notice check",
               7: "PART B — Approvals, settlement and closing"},
        outcome="Final settlement calculated · certificate issued · system access revoked",
        leads="— the journey is complete.",
    ),
]


# ---------------------------------------------------------------- helpers ---
def set_cell_bg(cell, hexc):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), hexc); pr.append(shd)


def page_num(p):
    r = p.add_run()
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
    b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = 'PAGE'
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'end')
    r._r.append(a); r._r.append(b); r._r.append(c)


def style_doc(doc):
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    for sz, nm in [(17, 'Heading 1'), (14, 'Heading 2'), (12, 'Heading 3')]:
        st = doc.styles[nm]; st.font.name = 'Calibri'; st.font.size = Pt(sz); st.font.color.rgb = NAVY
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.8)
        s.left_margin = s.right_margin = Inches(0.85)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(10)
        set_cell_bg(c, '1F4E79')
    for i, row in enumerate(rows):
        cs = t.add_row().cells
        for j, v in enumerate(row):
            cs[j].text = str(v)
            for p in cs[j].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9.5)
        if i % 2 == 1:
            for c in cs:
                set_cell_bg(c, 'F2F2F2')
    if widths:
        for r in t.rows:
            for j, w in enumerate(widths):
                r.cells[j].width = Inches(w)
    return t


MISSING = []


def embed(doc, path, width=6.3):
    p = Path(path)
    if not p.exists():
        MISSING.append(p.name)
        q = doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = q.add_run(f"[missing: {p.name}]"); r.font.color.rgb = RED
        return
    ip = doc.add_paragraph(); ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.add_run().add_picture(str(p), width=Inches(width))


def kv(doc, label, text, colour):
    p = doc.add_paragraph()
    r = p.add_run(f"{label}  "); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = colour
    r2 = p.add_run(text); r2.font.size = Pt(10.5)
    return p


# ------------------------------------------------------------------ build ---
def build():
    doc = Document(); style_doc(doc)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = HEADER_TEXT
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size = Pt(9); r.italic = True; r.font.color.rgb = GREY
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ").font.size = Pt(9); page_num(fp)

    # title
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Mukti Cox's Bazar"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = NAVY
    s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(DOC_TITLE); r.bold = True; r.font.size = Pt(21)
    s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s2.add_run(DOC_SUB)
    r.italic = True; r.font.size = Pt(13); r.font.color.rgb = GREY
    doc.add_page_break()

    # how to read
    doc.add_heading("How to read this guide", level=1)
    doc.add_paragraph(
        "Every step below is a real screen from the system. The button you must click "
        "is ringed in red, with a white pointer next to it and the step number beside it. "
        "The dark blue bar under each picture repeats the instruction, so the same picture "
        "can be dropped straight into a video.")
    table(doc, ["On the picture", "What it means"], [
        ("Red ring + pointer", "Click exactly here."),
        ("Red circle with a number", "Which step of the flow you are on."),
        ("Yellow badge, top right", "Who is doing this — the employee, or the approver."),
        ("Dark blue bar at the bottom", "The instruction for this step."),
    ], widths=[1.9, 4.5])
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Note about the two logins:  "); r.bold = True
    p.add_run("an approval flow needs two different people. The screenshots were taken "
              "by logging in as the employee for the first half, then logging in as the "
              "manager for the second half. Do the same when you record the video — it is "
              "the only way to show the hand-over honestly.")

    # contents / how the flows chain together
    doc.add_page_break()
    doc.add_heading(f"The {INTRO_COUNT} flows, and how they connect", level=1)
    doc.add_paragraph(
        "Each flow below produces something the next one needs. Read them in this "
        "order the first time — after that, jump straight to the one you need.")
    rows = []
    for i, fd in enumerate(FLOW_DEFS, start=1):
        src = FLOWS / fd["key"]
        n = len(json.loads((src / "steps.json").read_text())) if (src / "steps.json").exists() else 0
        rows.append([str(i), fd["title"].split("— ", 1)[-1], f"{n} steps",
                     fd["outcome"], fd.get("leads", "")])
    table(doc, ["#", "Flow", "Steps", "What it produces", "Which feeds"], rows,
          widths=[0.35, 1.7, 0.6, 2.1, 1.8])

    # each flow
    for fd in FLOW_DEFS:
        src = FLOWS / fd["key"]
        if not (src / "steps.json").exists():
            print(f"  !! no capture for '{fd['key']}' - skipped")
            continue
        steps = json.loads((src / "steps.json").read_text())

        doc.add_page_break()
        doc.add_heading(fd["title"], level=1)
        doc.add_paragraph(fd["summary"])

        doc.add_heading("Who does what", level=2)
        table(doc, ["Who", "Does what", "When"], fd["who"], widths=[2.0, 3.0, 1.4])

        doc.add_heading("Logins used", level=2)
        table(doc, ["Role", "Username", "Password"], fd["logins"], widths=[2.2, 2.1, 2.1])

        doc.add_heading("The flow at a glance", level=2)
        table(doc, ["Step", "Screen", "You click"],
              [[str(s["n"]), s["headline"], s["instruction"]] for s in steps],
              widths=[0.5, 2.3, 3.8])

        doc.add_heading("Step by step", level=2)
        for s in steps:
            if s["n"] in fd["parts"]:
                hp2 = doc.add_paragraph()
                r = hp2.add_run(fd["parts"][s["n"]])
                r.bold = True; r.font.size = Pt(13); r.font.color.rgb = GREEN

            h = doc.add_paragraph()
            r = h.add_run(f"Step {s['n']} — {s['headline']}")
            r.bold = True; r.font.size = Pt(12.5); r.font.color.rgb = NAVY
            embed(doc, src / s["file"])
            kv(doc, "WHO:", s["role"] or "—", GREEN)
            kv(doc, "YOU CLICK:", s["instruction"], RED)
            nxt = next((x for x in steps if x["n"] == s["n"] + 1), None)
            kv(doc, "THEN YOU SEE:", nxt["headline"] if nxt else fd["outcome"], NAVY)
            if s.get("note"):
                kv(doc, "TIP:", s["note"], GREY)
            doc.add_paragraph()

        doc.add_heading("Result", level=2)
        doc.add_paragraph(fd["outcome"])
        if fd.get("leads"):
            p = doc.add_paragraph()
            r = p.add_run("Leads to:  "); r.bold = True; r.font.color.rgb = GREEN
            p.add_run(fd["leads"])

    # recording notes
    doc.add_page_break()
    doc.add_heading("Recording these as videos", level=1)
    table(doc, ["Setting", "Use"], [
        ("Frame", "Keep the picture perfectly still. Do NOT add zoom or pan on a screen "
                  "recording — it makes the text shimmer and look like the screen is shaking."),
        ("Pace", "Hold each screen about 5 seconds. Pause 1 second before every click."),
        ("Cursor", "Turn on click-highlighting in your recorder."),
        ("Two logins", "Record the employee half, log out, record the approver half."),
        ("Resolution", "1920 × 1080, browser at 100 % zoom, full screen (F11)."),
    ], widths=[1.3, 5.1])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")
    d2 = Document(str(OUT))
    print(f"  images: {len(d2.inline_shapes)} · tables: {len(d2.tables)}")
    print("  missing:", MISSING or "none")


if __name__ == "__main__":
    build()
