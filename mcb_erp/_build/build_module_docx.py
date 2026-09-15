# -*- coding: utf-8 -*-
"""Build a Bangla trainer's guide for a module, in the same house style as the HR one."""
import json
import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, "/Data/odoo19_enterprise/custom-addons/mcb_hr/_build")
sys.path.insert(0, "/Data/odoo19_enterprise/custom-addons/mcb_erp/_build")
import build_training_docx as B     # noqa: E402  (house style helpers)

FLOWS = Path("/Data/odoo19_enterprise/custom-addons/mcb_hr/_build/flows")
URL = "https://crevice-flop-copious.ngrok-free.dev"


def build(cfg):
    sessions = cfg["sessions"]
    doc = Document()
    B.style_doc(doc)
    B.bn_run(doc.sections[0].header.paragraphs[0].add_run(cfg["header"]), 8.5, False, B.GREY)

    # ---- cover ----
    if B.LOGO.exists():
        lp = doc.add_paragraph(); lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.add_run().add_picture(str(B.LOGO), width=B.Inches(1.3))
    B.para(doc, "মুক্তি কক্সবাজার", 19, True, B.NAVY, 2, WD_ALIGN_PARAGRAPH.CENTER)
    B.para(doc, cfg["title_bn"], 25, True, B.NAVY, 4, WD_ALIGN_PARAGRAPH.CENTER)
    B.para(doc, cfg["title_en"], 12, False, B.GREY, 16, WD_ALIGN_PARAGRAPH.CENTER)
    B.para(doc, cfg["subtitle"], 13, False, None, 4, WD_ALIGN_PARAGRAPH.CENTER)
    B.para(doc, cfg["story_line"], 11, False, B.GREY, 24, WD_ALIGN_PARAGRAPH.CENTER)
    total_min = sum(s["minutes"] for s in sessions)
    total_steps = sum(len(s["steps"]) for s in sessions)
    B.table(doc, ["", ""], [
        ["সেশন সংখ্যা", f"{B.bn_num(len(sessions))} টি"],
        ["মোট সময়", f"আনুমানিক {B.bn_num(total_min)} মিনিট"],
        ["মোট ধাপ", f"{B.bn_num(total_steps)} টি স্ক্রিন"],
        ["সিস্টেমের ঠিকানা", URL],
        ["লগইন", "admin / admin"],
        ["সংস্করণ", "১.০ · সেপ্টেম্বর ২০২৬"],
    ], widths=[1.9, 4.5])
    B.page_break(doc)

    # ---- how to use ----
    B.heading(doc, "কীভাবে এই গাইডটি ব্যবহার করবেন")
    B.para(doc, "এটি প্রশিক্ষকের গাইড — অংশগ্রহণকারীদের হাতে দেওয়ার জন্য নয়। প্রতিটি "
                "ধাপে স্ক্রিনের ছবি, আর ছবির নিচে চারটি লাইন:")
    B.kv(doc, "বলুন", "এটুকু মুখে বলবেন। হুবহু পড়ার দরকার নেই, ভাব ঠিক রেখে নিজের ভাষায় বলুন।", B.NAVY)
    B.kv(doc, "করুন", "স্ক্রিনে ঠিক এই জায়গায় ক্লিক করবেন। ইংরেজিতে, কারণ স্ক্রিনেও ইংরেজিতেই আছে।", B.RED)
    B.kv(doc, "দেখাবে", "ক্লিকের পর যা আসবে।", B.GREEN)
    B.kv(doc, "খেয়াল রাখুন", "এখানেই মানুষ সবচেয়ে বেশি ভুল করে। এই লাইনটা বাদ দেবেন না।", B.AMBER)
    B.para(doc, "")
    B.para(doc, "একটি নিয়ম সবচেয়ে জরুরি", 12, True, B.NAVY, 3)
    B.para(doc, "বাটন বা মেনুর নাম কখনও বাংলায় অনুবাদ করবেন না। স্ক্রিনে ইংরেজিতেই "
                "লেখা থাকবে — অনুবাদ করে বললে তাঁরা খুঁজে পাবেন না। বাকি সব কথা বাংলায়।")
    B.para(doc, "")
    B.para(doc, "প্রস্তুতির তালিকা", 12, True, B.NAVY, 3)
    for item in [
        "প্রজেক্টর ঠিক আছে কি না দেখে নিন — পেছনের সারি থেকেও যেন পড়া যায়।",
        f"ঠিকানাটা আগেই ব্রাউজারে খুলে রাখুন: {URL}",
        "আগের দিন নিজে একবার পুরো পথটা হেঁটে দেখুন।",
        "প্রত্যেক অংশগ্রহণকারীর জন্য আলাদা লগইন রাখুন।",
        "হাতে-কলমে অনুশীলনের জন্য অন্তত ৩০ মিনিট আলাদা রাখুন।",
    ]:
        p = doc.add_paragraph(style="List Number"); B.bn_run(p.add_run(item), 10.5)
    B.page_break(doc)

    # ---- policy ----
    B.heading(doc, "নীতিমালা ও নিয়ম — এক নজরে")
    B.para(doc, "প্রশিক্ষণে সংখ্যা নিয়ে প্রশ্ন আসবেই। নিচের তথ্য সিস্টেমের কোড থেকে "
                "যাচাই করা — নির্ভয়ে বলতে পারেন।", 10.5, False, B.GREY)
    for title, headers, rows, widths in cfg["policy"]:
        B.para(doc, "")
        B.para(doc, title, 12, True, B.NAVY, 4)
        B.table(doc, headers, rows, widths=widths, sizes=(9, 9))
    B.page_break(doc)

    # ---- agenda ----
    B.heading(doc, "প্রশিক্ষণ সূচি")
    if cfg.get("parts"):
        B.table(doc, ["", "বিষয়", "পর্ব", "কত ঘন ঘন", "সময়"],
                [[f"সেশন {B.bn_num(s['no'])}", s["bn"],
                  cfg["parts"].get(s.get("part"), ("", "", ""))[0].split("—")[0].strip(),
                  s.get("freq", ""), f"{B.bn_num(s['minutes'])} মিনিট"]
                 for s in sessions],
                widths=[0.8, 1.9, 1.2, 1.4, 0.9], sizes=(9, 8.5))
    else:
        B.table(doc, ["", "বিষয়", "কারা থাকবেন", "সময়"],
                [[f"সেশন {B.bn_num(s['no'])}", s["bn"], s["who"], f"{B.bn_num(s['minutes'])} মিনিট"]
                 for s in sessions],
                widths=[0.85, 2.1, 2.5, 0.95], sizes=(9.5, 9))
    B.para(doc, "")
    B.para(doc, cfg["agenda_note"], 10.5, False, B.GREY)
    B.page_break(doc)

    # ---- sessions ----
    parts = cfg.get("parts") or {}
    seen_parts = set()
    for s in sessions:
        part = s.get("part")
        if part and part in parts and part not in seen_parts:
            seen_parts.add(part)
            title, blurb, who = parts[part]
            B.para(doc, "")
            B.para(doc, title, 22, True, B.NAVY, 6, WD_ALIGN_PARAGRAPH.CENTER)
            B.para(doc, blurb, 12, False, None, 10, WD_ALIGN_PARAGRAPH.CENTER)
            B.para(doc, who, 11, True, B.GREY, 14, WD_ALIGN_PARAGRAPH.CENTER)
            B.page_break(doc)
        B.heading(doc, f"সেশন {B.bn_num(s['no'])} — {s['bn']}")
        B.para(doc, s["en"], 11, False, B.GREY, 8)
        if s.get("freq"):
            B.table(doc, ["কারা থাকবেন", "কত ঘন ঘন", "সময়", "ভিডিও"],
                    [[s["who"], s["freq"], f"{B.bn_num(s['minutes'])} মিনিট", s["video"]]],
                    widths=[2.3, 1.5, 0.9, 1.7], sizes=(9, 9))
        else:
            B.table(doc, ["কারা থাকবেন", "সময়", "ভিডিও"],
                    [[s["who"], f"{B.bn_num(s['minutes'])} মিনিট", s["video"]]],
                    widths=[2.8, 1.1, 2.5], sizes=(9, 9))
        B.para(doc, "")
        B.para(doc, "গল্পটা", 12, True, B.NAVY, 3)
        B.para(doc, s["story"], 11)
        B.para(doc, "")
        B.para(doc, "এই সেশনে যা শেখানো হবে", 12, True, B.NAVY, 3)
        for g in s["goals"]:
            p = doc.add_paragraph(style="List Bullet"); B.bn_run(p.add_run(g), 10.5)
        B.para(doc, "")
        B.para(doc, "ধাপে ধাপে", 12, True, B.NAVY, 6)

        sf = FLOWS / s["key"] / "steps.json"
        steps = json.loads(sf.read_text()) if sf.exists() else []
        used = set()
        for st in steps:
            slug = st["slug"]
            key = slug if slug not in used else f"{slug}_2"
            used.add(slug)
            n = s["steps"].get(key) or s["steps"].get(slug)
            if not n:
                continue
            hp = doc.add_paragraph()
            hp.paragraph_format.space_before = B.Pt(10)
            hp.paragraph_format.space_after = B.Pt(4)
            hp.paragraph_format.keep_with_next = True
            B.bn_run(hp.add_run(f"ধাপ {B.bn_num(st['n'])}  ·  "), 11, True, B.NAVY)
            B.bn_run(hp.add_run(st["headline"]), 11, True, B.GREY)
            B.embed(doc, FLOWS / s["key"] / st["file"])
            B.kv(doc, "বলুন", n["say"], B.NAVY)
            B.kv(doc, "করুন", n["do"], B.RED)
            B.kv(doc, "দেখাবে", n["see"], B.GREEN)
            if n.get("tip"):
                B.kv(doc, "খেয়াল রাখুন", n["tip"], B.AMBER)

        B.para(doc, "")
        B.para(doc, "যেখানে সবচেয়ে বেশি ভুল হয়", 12, True, B.RED, 3)
        for m in s["mistakes"]:
            p = doc.add_paragraph(style="List Bullet"); B.bn_run(p.add_run(m), 10.5)
        B.para(doc, "")
        B.para(doc, "যাচাই প্রশ্ন", 12, True, B.NAVY, 3)
        for q in s["questions"]:
            p = doc.add_paragraph(style="List Number"); B.bn_run(p.add_run(q), 10.5)
        B.para(doc, "")
        B.para(doc, "হাতে-কলমে অনুশীলন", 12, True, B.GREEN, 3)
        B.para(doc, s["practice"], 10.5)
        B.page_break(doc)

    # ---- quick reference + faq ----
    B.heading(doc, "দ্রুত রেফারেন্স — কর্মীদের হাতে দেওয়ার জন্য")
    B.para(doc, "এই পাতাটা আলাদা প্রিন্ট করে দিয়ে দিতে পারেন।", 10.5, False, B.GREY)
    B.para(doc, "")
    B.table(doc, ["যা করতে চান", "কোথায় যাবেন"], cfg["quickref"], widths=[2.6, 3.8])
    B.page_break(doc)
    B.heading(doc, "সাধারণ প্রশ্ন ও উত্তর")
    for q, a in cfg["faq"]:
        B.para(doc, "")
        B.para(doc, f"প্রশ্ন: {q}", 11, True, B.NAVY, 2)
        B.para(doc, f"উত্তর: {a}", 10.5)

    out = Path(cfg["out"])
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    narrated = sum(1 for s in sessions
                   for st in (json.loads((FLOWS / s['key'] / 'steps.json').read_text())
                              if (FLOWS / s['key'] / 'steps.json').exists() else []))
    print(f"Wrote {out} ({out.stat().st_size:,} bytes)")
    print(f"  sessions: {len(sessions)} · screens: {narrated} · missing images: {B.MISSING or 'none'}")
