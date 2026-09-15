"""Build MCB_Projects_Module_Guide.docx — SRS v3 Part D (Projects).
Covers mcb_project + mcb_volunteer + mcb_vehicle + mcb_store.
Run: /Data/odoo19_enterprise/venv/bin/python _build/build_docx_projects.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path("/Data/odoo19_enterprise/custom-addons/mcb_erp/_build")
SHOT = ROOT / "screenshots" / "projects"
OUT = Path("/Data/odoo19_enterprise/mukticox/MCB_Projects_Module_Guide.docx")

# =============================================================================
# REQUIREMENT INVENTORY — SRS v3 Part D §22-29 (Projects)
# =============================================================================
REQS = [
    ("PRJ-001", "Part D", "Must", "Project profile: donor, contract number, total budget (BDT + donor currency), implementation area, activity code, linked Annex-27 budget",
     "Done", "project.project extension — MCB Profile tab"),
    ("PRJ-005", "Part D", "Must", "Staff-to-project % time allocation per period; an employee can never be allocated over 100% for an overlapping period",
     "Done", "mcb.project.assignment + @api.constrains _check_total"),
    ("PRJ-006", "Part D", "Should", "Project close-out checklist (final report / assets handed over / budget reconciled) before a project can be marked closed",
     "Done", "project.project: mcb_closeout_* booleans + action_mcb_closeout()"),
    ("PRJ-007", "Part D", "Should", "Chief Executive can lock a project's budget at close-out",
     "Done", "shared with BUD-008 (mcb_budget) — budget.analytic.action_mcb_lock()"),
    ("TMS-001", "Part D", "Must", "Individual Timesheet (Annexure-23) per staff member for a given month",
     "Done", "mcb.timesheet.report.wizard"),
    ("TMS-002", "Part D", "Must", "Timesheet shows an 'At a Glance' policy-vs-actual split across the employee's assigned projects",
     "Done", "mcb.timesheet.report.wizard._rows() cross-references mcb.project.assignment"),
    ("TMS-003", "Part D", "Should", "Draft timesheet lines auto-created overnight from the previous day's attendance, split by each employee's assignment %",
     "Done", "mcb.project.assignment._cron_autofill_timesheets()"),
    ("TMS-004", "Part D", "Must", "Monthly payroll cost allocation posts one journal entry spreading each employee's wage across their assigned projects/donors by %",
     "Done", "mcb.payroll.allocation.wizard.action_post_allocation()"),
    ("TMS-005", "Part D", "Must", "Timesheets must be validated by a supervisor before they count as final",
     "Done", "native Odoo timesheet_grid validation"),
    ("MON-003", "Part D", "Should", "Project Manager auto-alerted when a milestone slips past its deadline",
     "Done", "project.milestone cron"),
    ("MON-004", "Part D", "Must", "Quarterly project report: planned vs completed activities, % achievement, beneficiary reach, challenges & next-quarter plan",
     "Done", "mcb.quarterly.report.wizard"),
    ("MON-005", "Part D", "Must", "Beneficiary reach recorded per activity per period, disaggregated by gender, age band and camp/host location",
     "Done", "mcb.beneficiary.entry"),
    ("MON-006", "Part D", "Must", "Employee Travel Authorization (Annexure-30) workflow with cost breakdown and multi-level approval",
     "Done", "mcb.travel.authorization: submit → finance → recommend → approve"),
    ("MON-007", "Part D", "Should", "MEAL indicator framework: target vs achieved vs % per project activity",
     "Partial", "mcb.meal.indicator scaffolded (target/achieved/%); full indicator framework awaits the MEAL officer's complete list"),
    ("ATT-001", "Part D", "Must", "Monthly Staff Attendance Information (Annexure-20) report",
     "Done", "mcb.attendance.report.wizard"),
    ("ATT-004", "Part D", "Should", "Camp-based staff attendance can be bulk-imported from a CSV sheet",
     "Done", "mcb.attendance.import.wizard"),
    ("ATT-006", "Part D", "Should", "Absenteeism analysis dashboard",
     "Partial", "Attendance data and pivot views are available; a bespoke absenteeism dashboard has not been built"),
    ("VOL-001", "Part D §26", "Must", "HTV / RTV / Community Volunteer database with camp, block and FCN/registration number",
     "Done", "mcb.volunteer"),
    ("ATT-005", "Part D §26", "Must", "Volunteer daily attendance, tracked separately per volunteer type",
     "Done", "mcb.volunteer.attendance"),
    ("VOL-002", "Part D §26", "Must", "Monthly incentive batch computed from attendance days × a configurable daily rate per volunteer category",
     "Done", "mcb.volunteer.incentive.batch.action_compute()"),
    ("VOL-004", "Part D §26", "Must", "Incentive amount = attendance days × category daily rate (HTV/RTV/Community rates configurable per batch)",
     "Done", "mcb.volunteer.incentive.line._compute_amount()"),
    ("VOL-007", "Part D §26", "Must", "Paying a batch posts one journal entry: Dr incentive expense (by project analytic) / Cr cash-bank",
     "Done", "mcb.volunteer.incentive.batch.action_pay()"),
    ("VEH-001", "Part D §28", "Must", "Vehicle/Motorcycle Log Book (Annexure-10) daily trip line; saving a line syncs the fleet odometer reading",
     "Done", "mcb.vehicle.log"),
    ("VEH-002", "Part D §28", "Must", "Insurance & fitness-certificate expiry dates tracked per vehicle with a 30-day-ahead alert",
     "Done", "fleet.vehicle: mcb_insurance_expiry / mcb_fitness_expiry + cron"),
    ("VEH-003", "Part D §28", "Should", "Monthly Log Book usage summary, PDF + XLSX",
     "Done", "mcb.logbook.wizard"),
    ("MOV-001", "Part D §28", "Must", "Daily Movement Register (Annexure-09): who travelled, where, when, for what purpose",
     "Done", "mcb.movement.register"),
    ("MOV-004", "Part D §28", "Should", "Movement Register entry cross-linked to its originating Travel Authorization",
     "Done", "mcb.movement.register.travel_auth_id"),
    ("VEH-005", "Part D §28", "Should", "Automatic cross-reference between the vehicle log book and the movement register for the same trip",
     "Partial", "Both registers exist and can reference the same travel authorization; automatic line-level cross-matching is still manual"),
    ("STO-001", "Part D §29", "Must", "Store Register (Annexure-08): running balance per item from native stock moves",
     "Done", "mcb.store.register.wizard.action_print_register()"),
    ("STO-002", "Part D §29", "Should", "Monthly Stock Report: opening / receipts / issues / closing per item",
     "Done", "mcb.store.register.wizard.action_print_monthly()"),
    ("STO-004", "Part D §29", "Must", "Store Requisition Form (SRF): request → approval → internal-transfer issue; the resulting stock moves ARE the register entries",
     "Done", "mcb.srf: action_request/action_approve/action_issue()"),
    ("STO-007", "Part D §29", "Must", "NFI distribution muster roll with beneficiary names/IDs, linked to the issuing transfer",
     "Done", "mcb.muster.roll + mcb.muster.roll.line"),
]


# =============================================================================
# DOCX HELPERS (shared style with the HR / Purchase / Accounts guides)
# =============================================================================

def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)


def add_page_number(footer_para):
    run = footer_para.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def style_heading(doc):
    for size, name in [(16, 'Heading 1'), (14, 'Heading 2'), (12, 'Heading 3')]:
        st = doc.styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x36, 0x4D)


def set_default_font(doc):
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def embed(doc, path, cap, width_inches=6.0):
    p = path if isinstance(path, Path) else Path(path)
    if not p.exists():
        print(f"  MISSING: {p}")
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(f"[Screenshot missing: {p.name}]")
        r.font.color.rgb = RGBColor(0xAA, 0x00, 0x00)
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_img.add_run()
    r.add_picture(str(p), width=Inches(width_inches))
    caption(doc, cap)


def make_table(doc, headers, rows, banded=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_bg(cell, '1F4E79')
    for idx, row in enumerate(rows):
        rcells = table.add_row().cells
        for i, val in enumerate(row):
            rcells[i].text = str(val) if val is not None else ""
            for p in rcells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
        if banded and idx % 2 == 1:
            for c in rcells:
                set_cell_bg(c, 'F2F2F2')
    return table


def set_margins(doc):
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)


def status_symbol(s):
    return {"Done": "✅ Done", "Partial": "⚠ Partial", "Deferred": "❌ Deferred"}.get(s, s)


# =============================================================================
# BUILD
# =============================================================================

def build():
    doc = Document()
    set_default_font(doc)
    style_heading(doc)
    set_margins(doc)

    section = doc.sections[0]
    hp = section.header.paragraphs[0]
    hp.text = "Mukti Cox's Bazar — Projects Module on Odoo 19 Enterprise"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.size = Pt(9)
        run.italic = True
        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.add_run("Page ").font.size = Pt(9)
    add_page_number(fp)

    # ---- TITLE PAGE ----
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(4):
        p.add_run("\n")
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Mukti Cox's Bazar")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor(0x1F, 0x36, 0x4D)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Projects Module on Odoo 19 Enterprise")
    r.bold = True
    r.font.size = Pt(20)
    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub2.add_run("A Plain-English User Guide — SRS v3, Part D (Projects, Volunteers, Vehicles & Store)")
    r.italic = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph()
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Version 1.0  ·  9 July 2026  ·  MCB ERP Team").font.size = Pt(12)
    doc.add_paragraph()
    doc.add_paragraph()
    classif = doc.add_paragraph()
    classif.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = classif.add_run("CONFIDENTIAL — Internal Use Only")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0xAA, 0x00, 0x00)
    doc.add_page_break()

    # ---- TOC ----
    doc.add_heading("Table of Contents", level=1)
    toc_rows = [
        ("0", "Executive Summary — who this guide is for"),
        ("1", "Architecture Overview — the 4 modules behind 'Projects'"),
        ("2", "Chapter 1 — Project Profile (MCB Tab)"),
        ("3", "Chapter 2 — Staff Assignments (% Time)"),
        ("4", "Chapter 3 — Beneficiary Data & MEAL Indicators"),
        ("5", "Chapter 4 — Travel Authorization (Annexure-30)"),
        ("6", "Chapter 5 — Quarterly, Timesheet & Attendance Reports"),
        ("7", "Chapter 6 — Payroll Cost Allocation"),
        ("8", "Chapter 7 — Volunteer Database"),
        ("9", "Chapter 8 — Volunteer Incentive Batch"),
        ("10", "Chapter 9 — Vehicle Log Book & Movement Register"),
        ("11", "Chapter 10 — Store (SRF, Muster Roll, Store Register)"),
        ("12", "Consolidated SRS Traceability Matrix"),
        ("A", "Appendix A — Glossary of Odoo Terms"),
    ]
    tbl = doc.add_table(rows=len(toc_rows), cols=2)
    tbl.style = "Light List Accent 1"
    for i, (ch, name) in enumerate(toc_rows):
        cells = tbl.rows[i].cells
        cells[0].text = ch
        cells[1].text = name
        for cell in cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(11)
    doc.add_page_break()

    # ---- CHAPTER 0: EXEC SUMMARY ----
    doc.add_heading("0. Executive Summary", level=1)
    doc.add_paragraph(
        "This guide is written for Project Managers, MEAL officers, HR/Finance "
        "staff who process staff time and travel, and anyone running the "
        "Volunteer, Vehicle or Store programmes at Mukti Cox's Bazar (MCB) — no "
        "assumption of prior software experience is made. It covers everything "
        "under SRS v3 Part D, with a screenshot for every screen."
    )
    doc.add_paragraph(
        "Four modules work together here. This guide is organised by everyday "
        "task, but every chapter states the module and the exact SRS requirement "
        "it satisfies."
    )
    doc.add_heading("Scope & Coverage", level=2)
    total = len(REQS)
    done = sum(1 for r in REQS if r[4] == "Done")
    partial = sum(1 for r in REQS if r[4] == "Partial")
    make_table(doc, ["Metric", "Value"], [
        ("Requirements indexed (SRS v3, Part D §22–29)", f"{total}"),
        ("✅ Done", f"{done}"),
        ("⚠ Partial", f"{partial}"),
        ("Odoo modules", "mcb_project, mcb_volunteer, mcb_vehicle, mcb_store"),
    ])
    doc.add_paragraph(
        "Everything in this guide was captured directly from a live, running copy "
        "of the system using real (demo) data driven through the actual approval "
        "buttons — nothing here is a mock-up. Three items remain Partial: MON-007 "
        "(full MEAL indicator framework awaits the MEAL officer's complete "
        "indicator list), ATT-006 (a bespoke absenteeism dashboard, though the "
        "underlying attendance data is fully available), and VEH-005 (automatic "
        "cross-referencing between the vehicle log and movement register, "
        "currently a manual cross-check)."
    )
    doc.add_page_break()

    # ---- CHAPTER 1: ARCHITECTURE ----
    doc.add_heading("1. Architecture Overview", level=1)
    make_table(doc, ["Module", "Extends", "Covers"], [
        ("mcb_project", "project_enterprise, hr_timesheet, hr_attendance", "Project profile, staff assignments, beneficiaries, travel authorization, MEAL, quarterly/timesheet/attendance reports, payroll allocation"),
        ("mcb_volunteer", "(standalone MCB app)", "HTV/RTV/Community volunteer database, attendance, incentive batches"),
        ("mcb_vehicle", "fleet", "Vehicle log book, movement register, expiry alerts, monthly reports"),
        ("mcb_store", "stock", "Store Requisition Form, Store Register, NFI muster rolls"),
    ])
    embed(doc, SHOT / "00_overview/01_apps_mcb_filtered.png", "Figure 1.1 — MCB apps on the home dashboard")
    doc.add_page_break()

    def chapter_trace(prefix_list, chapter_no):
        rows = [r for r in REQS if r[0].startswith(tuple(prefix_list))]
        body = [[r[0], r[1], r[2], r[3], status_symbol(r[4]), r[5]] for r in rows]
        doc.add_heading(f"{chapter_no} SRS Traceability", level=3)
        make_table(doc, ["Req ID", "SRS §", "Priority", "Requirement", "Status", "Implementation"], body)

    # ====== CHAPTER 2: PROJECT PROFILE ======
    doc.add_heading("2. Project Profile (MCB Tab)", level=1)
    doc.add_paragraph(
        "Every donor-funded project in Odoo carries an 'MCB Profile' tab holding "
        "everything the SRS asks for: donor, contract number, budget in both BDT "
        "and donor currency, implementation area, and a close-out checklist that "
        "must be completed before the project can be locked."
    )
    doc.add_heading("2.1 Where to find it", level=2)
    doc.add_paragraph("Project → Projects.")
    embed(doc, SHOT / "01_project_profile/01_project_list.png", "Figure 2.1 — Projects list")
    doc.add_heading("2.2 Walk-through", level=2)
    embed(doc, SHOT / "01_project_profile/02_gbvie_project_form.png", "Figure 2.2 — GBViE project — Description tab")
    doc.add_paragraph(
        "The MCB Profile tab carries the donor/budget data (PRJ-001) and the "
        "close-out checklist (PRJ-006/007) — the 'Close Out Project' button only "
        "unlocks once all three boxes are ticked."
    )
    embed(doc, SHOT / "01_project_profile/03_gbvie_project_mcb_profile_tab.png", "Figure 2.3 — MCB Profile tab — donor, budget, close-out checklist")
    chapter_trace(["PRJ-"], "2.3")
    doc.add_page_break()

    # ====== CHAPTER 3: ASSIGNMENTS ======
    doc.add_heading("3. Staff Assignments (% Time)", level=1)
    doc.add_paragraph(
        "Every staff member's time is split across the projects they support, as "
        "a percentage. Odoo will refuse to save an assignment that would push an "
        "employee over 100% for any overlapping period — this is what keeps "
        "donor cost-sharing honest."
    )
    doc.add_heading("3.1 Where to find it", level=2)
    doc.add_paragraph("Project → MCB → Staff Assignments (% Time).")
    embed(doc, SHOT / "02_assignments/01_assignment_list_pct_time.png", "Figure 3.1 — Staff Assignments — % time per project")
    chapter_trace(["PRJ-005"], "3.2")
    doc.add_page_break()

    # ====== CHAPTER 4: BENEFICIARY / MEAL ======
    doc.add_heading("4. Beneficiary Data & MEAL Indicators", level=1)
    doc.add_paragraph(
        "Field teams log how many people were reached by each activity, broken "
        "down by gender, age band and whether they were camp or host-community "
        "members. MEAL indicators track target vs achieved progress alongside."
    )
    doc.add_heading("4.1 Where to find it", level=2)
    doc.add_paragraph("Project → MCB → Beneficiary Data / MEAL Indicators.")
    embed(doc, SHOT / "03_beneficiary_meal/01_beneficiary_list.png", "Figure 4.1 — Beneficiary Data list")
    doc.add_heading("4.2 Walk-through", level=2)
    embed(doc, SHOT / "03_beneficiary_meal/02_beneficiary1_form.png", "Figure 4.2 — Beneficiary entry — gender/age/location breakdown")
    embed(doc, SHOT / "03_beneficiary_meal/03_meal_indicator_list.png", "Figure 4.3 — MEAL Indicators — target vs achieved")
    chapter_trace(["MON-005", "MON-007"], "4.3")
    doc.add_page_break()

    # ====== CHAPTER 5: TRAVEL AUTH ======
    doc.add_heading("5. Travel Authorization (Annexure-30)", level=1)
    doc.add_paragraph(
        "Any staff travel is authorised in advance: destination, purpose, "
        "transport mode, and a full cost breakdown (travel/lodging/food/"
        "registration/other), routed through Finance review and Line Manager "
        "recommendation before Chief Executive approval."
    )
    doc.add_heading("5.1 Where to find it", level=2)
    doc.add_paragraph("Project → MCB → Travel Authorizations (Annex-30).")
    embed(doc, SHOT / "04_travel_auth/01_travel_auth_list.png", "Figure 5.1 — Travel Authorizations list")
    doc.add_heading("5.2 Walk-through", level=2)
    embed(doc, SHOT / "04_travel_auth/02_travel_auth1_form_approved.png", "Figure 5.2 — Travel Authorization — approved, Dhaka coordination meeting")
    embed(doc, SHOT / "04_travel_auth/03_travel_auth1_report_preview.png", "Figure 5.3 — Travel Authorization (Annexure-30) — PDF preview")
    chapter_trace(["MON-006"], "5.3")
    doc.add_page_break()

    # ====== CHAPTER 6: REPORTS ======
    doc.add_heading("6. Quarterly, Timesheet & Attendance Reports", level=1)
    doc.add_paragraph(
        "Three report wizards give managers, donors and HR the numbers they need "
        "without touching a spreadsheet."
    )
    doc.add_heading("6.1 Quarterly Report (MON-004)", level=2)
    embed(doc, SHOT / "05_reports/01_quarterly_report_wizard_form.png", "Figure 6.1 — Quarterly Report wizard")
    embed(doc, SHOT / "05_reports/02_quarterly_report_preview.png", "Figure 6.2 — Quarterly Report — preview")
    doc.add_heading("6.2 Individual Timesheet (Annexure-23)", level=2)
    embed(doc, SHOT / "05_reports/03_timesheet_report_wizard_form.png", "Figure 6.3 — Timesheet Report wizard")
    embed(doc, SHOT / "05_reports/04_timesheet_report_preview.png", "Figure 6.4 — Individual Timesheet — At-a-Glance preview")
    doc.add_heading("6.3 Monthly Attendance (Annexure-20)", level=2)
    embed(doc, SHOT / "05_reports/05_attendance_report_wizard_form.png", "Figure 6.5 — Attendance Report wizard")
    embed(doc, SHOT / "05_reports/06_attendance_report_preview.png", "Figure 6.6 — Monthly Attendance — preview")
    doc.add_heading("6.4 Camp Attendance Import", level=2)
    embed(doc, SHOT / "05_reports/07_camp_attendance_import_wizard_form.png", "Figure 6.7 — Camp Attendance CSV Import wizard")
    chapter_trace(["TMS-001", "TMS-002", "TMS-003", "MON-004", "ATT-001", "ATT-004", "ATT-006"], "6.5")
    doc.add_page_break()

    # ====== CHAPTER 7: PAYROLL ALLOCATION ======
    doc.add_heading("7. Payroll Cost Allocation", level=1)
    doc.add_paragraph(
        "Once a month, Finance runs this wizard to spread every assigned "
        "employee's wage across their project analytic accounts by the exact % "
        "recorded in Chapter 3 — so each donor's ledger shows its true staff "
        "cost, in one journal entry."
    )
    doc.add_heading("7.1 Where to find it", level=2)
    doc.add_paragraph("Project → MCB → Payroll Cost Allocation (TMS-004).")
    embed(doc, SHOT / "06_payroll_allocation/01_payroll_allocation_wizard_form.png", "Figure 7.1 — Payroll Allocation wizard")
    embed(doc, SHOT / "06_payroll_allocation/02_payroll_allocation_je_posted.png", "Figure 7.2 — Resulting journal entry — posted, split by project")
    chapter_trace(["TMS-004", "TMS-005"], "7.2")
    doc.add_page_break()

    # ====== CHAPTER 8: VOLUNTEER DB ======
    doc.add_heading("8. Volunteer Database", level=1)
    doc.add_paragraph(
        "MCB runs community volunteer programmes using three categories — Host "
        "Trained Volunteer (HTV), Rohingya Trained Volunteer (RTV), and general "
        "Community Volunteer — each with their own camp/block, contact details "
        "and preferred payment method (cash / bKash / bank)."
    )
    doc.add_heading("8.1 Where to find it", level=2)
    doc.add_paragraph("MCB Volunteers → Volunteers Database / Attendance.")
    embed(doc, SHOT / "07_volunteer/01_volunteer_list.png", "Figure 8.1 — Volunteers list")
    doc.add_heading("8.2 Walk-through", level=2)
    embed(doc, SHOT / "07_volunteer/02_volunteer1_form.png", "Figure 8.2 — Volunteer — RTV, Camp 12")
    embed(doc, SHOT / "07_volunteer/03_volunteer_attendance_list.png", "Figure 8.3 — Volunteer Attendance list")
    chapter_trace(["VOL-001", "ATT-005"], "8.3")
    doc.add_page_break()

    # ====== CHAPTER 9: INCENTIVE BATCH ======
    doc.add_heading("9. Volunteer Incentive Batch", level=1)
    doc.add_paragraph(
        "Once a month, an Incentive Batch computes what every active volunteer "
        "is owed — attendance days × their category's daily rate — and, once "
        "approved, posts a single payment journal entry."
    )
    doc.add_heading("9.1 Where to find it", level=2)
    doc.add_paragraph("MCB Volunteers → Incentive Batches.")
    embed(doc, SHOT / "08_incentive/01_incentive_batch_list.png", "Figure 9.1 — Incentive Batches list")
    doc.add_heading("9.2 Walk-through", level=2)
    embed(doc, SHOT / "08_incentive/02_incentive_batch2_form_paid.png", "Figure 9.2 — Incentive Batch — computed, approved, paid")
    embed(doc, SHOT / "08_incentive/03_incentive_signature_sheet_preview.png", "Figure 9.3 — Incentive signature sheet — PDF preview")
    chapter_trace(["VOL-002", "VOL-004", "VOL-007"], "9.3")
    doc.add_page_break()

    # ====== CHAPTER 10: VEHICLE ======
    doc.add_heading("10. Vehicle Log Book & Movement Register", level=1)
    doc.add_paragraph(
        "Every vehicle trip is logged with driver, purpose, odometer readings "
        "(which automatically update the vehicle's official odometer) and fuel "
        "cost. A separate Movement Register tracks staff journeys and can be "
        "cross-linked to the Travel Authorization that approved the trip."
    )
    doc.add_heading("10.1 Where to find it", level=2)
    doc.add_paragraph("Fleet → Vehicles, and Fleet → MCB → Vehicle Log Book / Movement Register.")
    embed(doc, SHOT / "09_vehicle/01_vehicle_list.png", "Figure 10.1 — Vehicles list")
    doc.add_heading("10.2 Walk-through", level=2)
    doc.add_paragraph("The Toyota Corolla assigned to the GBViE project, with insurance/fitness expiry dates tracked.")
    embed(doc, SHOT / "09_vehicle/02_vehicle5_corolla_form_gbvie.png", "Figure 10.2 — Vehicle — assigned to GBViE, expiry dates set")
    embed(doc, SHOT / "09_vehicle/03_vehicle_log_list.png", "Figure 10.3 — Vehicle Log Book — trip lines with odometer & fuel")
    embed(doc, SHOT / "09_vehicle/04_movement_register_list.png", "Figure 10.4 — Movement Register list")
    embed(doc, SHOT / "09_vehicle/05_movement1_form_linked_travel_auth.png", "Figure 10.5 — Movement Register entry — linked to Travel Authorization")
    embed(doc, SHOT / "09_vehicle/06_movement1_report_preview.png", "Figure 10.6 — Movement Register (Annexure-09) — PDF preview")
    embed(doc, SHOT / "09_vehicle/07_logbook_wizard_form.png", "Figure 10.7 — Monthly Log Book wizard")
    embed(doc, SHOT / "09_vehicle/08_logbook_monthly_report_preview.png", "Figure 10.8 — Monthly Log Book (Annexure-10) — PDF preview")
    chapter_trace(["VEH-", "MOV-"], "10.3")
    doc.add_page_break()

    # ====== CHAPTER 11: STORE ======
    doc.add_heading("11. Store (SRF, Muster Roll, Store Register)", level=1)
    doc.add_paragraph(
        "NFI (Non-Food Item) distributions run through a Store Requisition Form: "
        "requested, approved, then issued as a real internal stock transfer. "
        "Every issued item automatically becomes a Store Register entry — there "
        "is no separate bookkeeping step. Distributions to beneficiaries are "
        "recorded on a signed muster roll."
    )
    doc.add_heading("11.1 Where to find it", level=2)
    doc.add_paragraph("Inventory → MCB Store → Store Requisitions / NFI Muster Rolls / Store Register.")
    embed(doc, SHOT / "10_store/01_srf_list.png", "Figure 11.1 — Store Requisitions (SRF) list")
    doc.add_heading("11.2 Walk-through", level=2)
    doc.add_paragraph("An SRF issuing tarpaulins, hygiene kits and jerry cans to the Camp 12 distribution team.")
    embed(doc, SHOT / "10_store/02_srf1_form_issued.png", "Figure 11.2 — SRF — issued, internal transfer created")
    embed(doc, SHOT / "10_store/03_muster_roll_list.png", "Figure 11.3 — Muster Rolls list")
    embed(doc, SHOT / "10_store/04_muster_roll1_form_distributed.png", "Figure 11.4 — Muster Roll — beneficiaries recorded, distributed")
    embed(doc, SHOT / "10_store/05_muster_roll1_report_preview.png", "Figure 11.5 — Muster Roll — PDF preview")
    embed(doc, SHOT / "10_store/06_store_register_wizard_form.png", "Figure 11.6 — Store Register wizard")
    embed(doc, SHOT / "10_store/07_store_register_report_preview.png", "Figure 11.7 — Store Register (Annexure-08) — running balance preview")
    embed(doc, SHOT / "10_store/08_monthly_stock_report_preview.png", "Figure 11.8 — Monthly Stock Report — preview")
    chapter_trace(["STO-"], "11.3")
    doc.add_page_break()

    # ====== CHAPTER 12: MASTER MATRIX ======
    doc.add_heading("12. Consolidated SRS Traceability Matrix", level=1)
    doc.add_paragraph(
        "Every Part D (Projects, Volunteers, Vehicles & Store) requirement "
        "implemented across the 4 modules, in one place."
    )
    rows = [[r[0], r[2], r[3], status_symbol(r[4])] for r in REQS]
    make_table(doc, ["Req ID", "Priority", "Requirement", "Status"], rows)
    doc.add_page_break()

    # ====== APPENDIX A: GLOSSARY ======
    doc.add_heading("Appendix A — Glossary of Odoo Terms", level=1)
    glossary = [
        ("project.project / project.task", "Odoo's native project and activity models — MCB adds the donor/budget/close-out MCB Profile tab."),
        ("Analytic account", "The cost/donor dimension attached to every project — used across timesheets, payroll allocation and budget reports."),
        ("hr.attendance", "Native Odoo check-in/check-out record; drives both the Attendance report and the timesheet auto-draft cron."),
        ("account.analytic.line", "Odoo's technical name for a timesheet line."),
        ("Cron / ir.cron", "A scheduled background job — used for milestone alerts, timesheet auto-draft, and vehicle expiry alerts."),
        ("fleet.vehicle", "Odoo's native vehicle record — MCB adds project assignment, expiry dates and the log book."),
        ("stock.picking / stock.move", "Odoo's native internal-transfer and stock-movement models — the SRF's issue step creates these, and they ARE the Store Register."),
        ("Wizard / Transient model", "A pop-up form used once to run an action or print a report — Quarterly/Timesheet/Attendance/Payroll-Allocation/Logbook/Store-Register are all wizards."),
        ("QWeb report", "Odoo's PDF templating engine — used for every printed form in this guide."),
        ("Chatter", "The message/activity panel on the right of every record — full audit history."),
        ("Sequence", "Auto-incrementing reference generator — e.g. TAF-2627-001, VIB-2627-001, SRF-2627-001, reset every July per fiscal year."),
    ]
    make_table(doc, ["Term", "Meaning"], glossary)

    # ---- SAVE + VALIDATE ----
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"\nWrote {OUT}  ({OUT.stat().st_size:,} bytes)")

    reread = Document(str(OUT))
    headings = [p.text for p in reread.paragraphs if p.style.name.startswith("Heading 1")]
    print("Top-level chapters present:")
    for h in headings:
        print(f"  • {h}")
    missing = []
    for area_dir in SHOT.iterdir():
        if not area_dir.is_dir():
            continue
        for png in area_dir.glob("*.png"):
            if not png.exists():
                missing.append(str(png))
    if missing:
        print("MISSING SCREENSHOTS:", missing)
    found = False
    for tbl in reread.tables:
        first_cells = [c.text for c in tbl.rows[0].cells]
        if first_cells == ["Req ID", "Priority", "Requirement", "Status"]:
            data_rows = len(tbl.rows) - 1
            print(f"Master traceability matrix rows: {data_rows}  (expected {len(REQS)})")
            assert data_rows == len(REQS), "matrix row count mismatch"
            found = True
            break
    assert found, "master matrix not found in document"
    print("Validation: OK")


if __name__ == "__main__":
    build()
