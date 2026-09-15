# -*- coding: utf-8 -*-
"""Build MCB_HR_Payslip_Training_Bangla.docx — the detailed payslip guide."""
import json
import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.path.insert(0, "/Data/odoo19_enterprise/custom-addons/mcb_hr/_build")
import build_training_docx as B          # noqa: E402  (reuse the house style)
import payslip_bn as P                   # noqa: E402

FLOW = Path("/Data/odoo19_enterprise/custom-addons/mcb_hr/_build/flows/payslip")
OUT = Path("/Data/odoo19_enterprise/mukticox/MCB_HR_Payslip_Training_Bangla.docx")
URL = "https://crevice-flop-copious.ngrok-free.dev"


def build():
    doc = Document()
    B.style_doc(doc)
    hp = doc.sections[0].header.paragraphs[0]
    B.bn_run(hp.add_run("মুক্তি কক্সবাজার · Odoo ১৯ — পে-স্লিপ ও বেতন · প্রশিক্ষকের গাইড"),
             8.5, False, B.GREY)

    # cover
    if B.LOGO.exists():
        lp = doc.add_paragraph(); lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.add_run().add_picture(str(B.LOGO), width=B.Inches(1.3))
    B.para(doc, "মুক্তি কক্সবাজার", 19, True, B.NAVY, 2, WD_ALIGN_PARAGRAPH.CENTER)
    B.para(doc, "পে-স্লিপ ও মাসিক বেতন", 25, True, B.NAVY, 4, WD_ALIGN_PARAGRAPH.CENTER)
    B.para(doc, "Payslip & Monthly Payroll — Detailed Trainer's Guide", 12, False, B.GREY,
           16, WD_ALIGN_PARAGRAPH.CENTER)
    B.para(doc, "একটি সংখ্যা থেকে নয়টি লাইন — ধাপে ধাপে, ক্রম অনুযায়ী",
           13, False, None, 4, WD_ALIGN_PARAGRAPH.CENTER)
    B.para(doc, "উদাহরণ: রাশেদুল করিম · মাসিক Wage ৪২,০০০ টাকা · অগাস্ট ২০২৬",
           11, False, B.GREY, 22, WD_ALIGN_PARAGRAPH.CENTER)

    steps = json.loads((FLOW / "steps.json").read_text())
    B.table(doc, ["", ""], [
        ["কারা থাকবেন", "ফিন্যান্স ও পে-রোল দায়িত্বপ্রাপ্ত (এইচআর চাইলে থাকতে পারেন)"],
        ["সময়", "আনুমানিক ৭৫ মিনিট"],
        ["মোট ধাপ", f"{B.bn_num(len(steps))} টি স্ক্রিন, পাঁচটি পর্বে"],
        ["সিস্টেমের ঠিকানা", URL],
        ["লগইন", "admin / admin"],
    ], widths=[1.9, 4.5])
    B.page_break(doc)

    # why this guide + the whole calculation on one page
    B.heading(doc, "শুরুর আগে — পুরো হিসাবটা এক পাতায়")
    B.para(doc, "পে-স্লিপ কঠিন লাগে কারণ মানুষ মাঝখান থেকে পড়া শুরু করে। আসলে পুরোটা "
                "একটি সংখ্যা থেকে তৈরি — চুক্তির মাসিক Wage। নিচের ছকটাই পুরো সেশনের "
                "মূল কথা। প্রশিক্ষণ শুরুর আগে এটি বোর্ডে লিখে রাখুন, আর শেষেও একবার "
                "দেখান।", 11)
    B.para(doc, "")
    B.para(doc, "রাশেদুল করিম — অগাস্ট ২০২৬ (সাধারণ মাস)", 12, True, B.NAVY, 4)
    B.table(doc, ["#", "লাইন", "কীভাবে এলো", "টাকা"], P.CALC_AUG,
            widths=[0.4, 2.3, 2.5, 1.2], sizes=(9.5, 9.5))
    B.para(doc, "")
    B.para(doc, "একই ব্যক্তি — এপ্রিল ২০২৬ (বোনাসের মাস)", 12, True, B.NAVY, 4)
    B.table(doc, ["", "লাইন", "কীভাবে এলো", "টাকা"], P.CALC_APR,
            widths=[0.4, 2.3, 2.5, 1.2], sizes=(9.5, 9.5))
    B.para(doc, "")
    B.para(doc, "আয়করের স্ল্যাব (বার্ষিক আয়)", 12, True, B.NAVY, 4)
    B.table(doc, ["বার্ষিক আয়ের স্তর", "হার", "রাশেদুলের ক্ষেত্রে কর"], P.TAX,
            widths=[2.5, 1.0, 2.9], sizes=(9.5, 9.5))
    B.para(doc, "")
    B.para(doc, "তিনটি কথা বারবার বলবেন", 12, True, B.NAVY, 4)
    for t in [
        "Wage আর Basic এক নয়। Basic = Wage-এর ৬০%। সব ভাতা ও পিএফ Basic থেকে হিসাব হয়।",
        "পে-স্লিপে কোনো সংখ্যা কেউ হাতে বসায় না। ভুল হলে গোড়ায় ঠিক করতে হয়।",
        "প্রতিষ্ঠানের পিএফ অংশ কর্তন নয় — ওটা MCB-র খরচ, কর্মীর বেতন কমায় না।",
    ]:
        p = doc.add_paragraph(style="List Bullet"); B.bn_run(p.add_run(t), 11)
    B.page_break(doc)

    # the steps, grouped into parts
    for st in steps:
        slug, n = st["slug"], st["n"]
        if slug in P.PARTS:
            B.heading(doc, P.PARTS[slug], level=1)
        nar = P.STEPS.get(slug)
        if not nar:
            continue
        hp2 = doc.add_paragraph()
        hp2.paragraph_format.space_before = B.Pt(10)
        hp2.paragraph_format.space_after = B.Pt(4)
        hp2.paragraph_format.keep_with_next = True
        B.bn_run(hp2.add_run(f"ধাপ {B.bn_num(n)}  ·  "), 11, True, B.NAVY)
        B.bn_run(hp2.add_run(st["headline"]), 11, True, B.GREY)
        B.embed(doc, FLOW / st["file"])
        B.kv(doc, "বলুন", nar["say"], B.NAVY)
        B.kv(doc, "করুন", nar["do"], B.RED)
        B.kv(doc, "দেখাবে", nar["see"], B.GREEN)
        if nar.get("tip"):
            B.kv(doc, "খেয়াল রাখুন", nar["tip"], B.AMBER)

    # wrap-up
    B.page_break(doc)
    B.heading(doc, "যেখানে সবচেয়ে বেশি ভুল হয়")
    for m in P.MISTAKES:
        p = doc.add_paragraph(style="List Bullet"); B.bn_run(p.add_run(m), 10.5)
    B.para(doc, "")
    B.heading(doc, "যাচাই প্রশ্ন", level=2)
    for q in P.QUESTIONS:
        p = doc.add_paragraph(style="List Number"); B.bn_run(p.add_run(q), 10.5)
    B.para(doc, "")
    B.heading(doc, "হাতে-কলমে অনুশীলন", level=2)
    for line in P.PRACTICE.split("\n"):
        B.para(doc, line, 10.5)
    B.page_break(doc)
    B.heading(doc, "কর্মীরা যে প্রশ্নগুলো করেন")
    for q, a in P.FAQ:
        B.para(doc, "")
        B.para(doc, f"প্রশ্ন: {q}", 11, True, B.NAVY, 2)
        B.para(doc, f"উত্তর: {a}", 10.5)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")
    print(f"  steps: {len(steps)} · narrated: {sum(1 for s in steps if s['slug'] in P.STEPS)}")
    print(f"  missing images: {B.MISSING or 'none'}")


if __name__ == "__main__":
    build()
