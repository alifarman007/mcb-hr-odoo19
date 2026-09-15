# -*- coding: utf-8 -*-
"""Build MCB_HR_Training_Guide_Bangla.docx — the trainer's guide for the HR module.

The narration is Bangla (what the trainer says); every button and menu name stays
English, because that is what is on the screen the trainee is looking at.
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

sys.path.insert(0, "/Data/odoo19_enterprise/custom-addons/mcb_hr/_build")
from training_bn import SESSIONS          # noqa: E402
from training_facts import POLICY, BEFORE_YOU_TRAIN, FAQ, QUICKREF  # noqa: E402

FLOWS = Path("/Data/odoo19_enterprise/custom-addons/mcb_hr/_build/flows")
LOGO = Path(__file__).resolve().parent / "assets" / "mcb_logo.png"   # MCB logo, from the Joining Letter template
OUT = Path("/Data/odoo19_enterprise/mukticox/MCB_HR_Training_Guide_Bangla.docx")

URL = "https://crevice-flop-copious.ngrok-free.dev"

# Nirmala UI ships with Windows and renders Bengali correctly; on Linux/LibreOffice
# fontconfig substitutes a Bengali-capable face automatically.
BN = "Nirmala UI"
EN = "Calibri"

NAVY = RGBColor(0x1F, 0x36, 0x4D)
RED = RGBColor(0xC0, 0x28, 0x28)
GREEN = RGBColor(0x1E, 0x7A, 0x53)
AMBER = RGBColor(0xA9, 0x6A, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)
MISSING = []

_BN_DIGITS = str.maketrans("0123456789", "০১২৩৪৫৬৭৮৯")


def bn_num(value):
    """Bengali numerals — the document mixes generated counts with hand-written
    Bangla, and Western digits beside Bengali ones read as a typo."""
    return str(value).translate(_BN_DIGITS)


# ------------------------------------------------------------------ helpers --
def bn_run(run, size=11, bold=False, color=None):
    """Bengali is a complex script: Word picks the font from w:cs, and the size
    from w:szCs, not from the latin attributes. Set both or it renders in a
    fallback face at the wrong size."""
    run.font.name = BN
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), BN)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), str(int(size * 2)))
    rpr.append(szcs)
    if bold:
        bcs = OxmlElement("w:bCs")
        rpr.append(bcs)
    return run


def para(doc, text="", size=11, bold=False, color=None, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if text:
        bn_run(p.add_run(text), size, bold, color)
    return p


def set_cell_bg(cell, hexc):
    pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hexc)
    pr.append(shd)


def style_doc(doc):
    n = doc.styles["Normal"]
    n.font.name = BN
    n.font.size = Pt(11)
    for sz, nm in [(18, "Heading 1"), (14, "Heading 2"), (12, "Heading 3")]:
        st = doc.styles[nm]
        st.font.name = BN
        st.font.size = Pt(sz)
        st.font.color.rgb = NAVY
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.8)


def heading(doc, text, level=1, color=NAVY):
    h = doc.add_heading("", level=level)
    bn_run(h.add_run(text), 18 - (level - 1) * 4, True, color)
    return h


def table(doc, headers, rows, widths=None, sizes=(9.5, 9.5)):
    if headers and not any(headers):          # headerless block
        t = doc.add_table(rows=0, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for i, row in enumerate(rows):
            cs = t.add_row().cells
            for j, v in enumerate(row):
                cs[j].text = ""
                bn_run(cs[j].paragraphs[0].add_run(str(v)), sizes[1], j == 0)
            if i % 2 == 1:
                for c in cs:
                    set_cell_bg(c, "F4F6F8")
        if widths:
            for r in t.rows:
                for j, w in enumerate(widths):
                    r.cells[j].width = Inches(w)
        return t
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        bn_run(c.paragraphs[0].add_run(h), sizes[0], True, RGBColor(255, 255, 255))
        set_cell_bg(c, "1F364D")
    for i, row in enumerate(rows):
        cs = t.add_row().cells
        for j, v in enumerate(row):
            cs[j].text = ""
            bn_run(cs[j].paragraphs[0].add_run(str(v)), sizes[1])
        if i % 2 == 1:
            for c in cs:
                set_cell_bg(c, "F4F6F8")
    if widths:
        for r in t.rows:
            for j, w in enumerate(widths):
                r.cells[j].width = Inches(w)
    return t


def embed(doc, path, width=6.4):
    p = Path(path)
    if not p.exists():
        MISSING.append(p.name)
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        bn_run(q.add_run(f"[ছবি নেই: {p.name}]"), 9, False, RED)
        return
    ip = doc.add_paragraph()
    ip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ip.paragraph_format.space_after = Pt(4)
    ip.add_run().add_picture(str(p), width=Inches(width))


def kv(doc, label, text, colour, size=10.5):
    """One narration line: a coloured Bangla label, then the content."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.12)
    bn_run(p.add_run(f"{label}  "), size, True, colour)
    bn_run(p.add_run(text), size)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# -------------------------------------------------------------------- build --
def build():
    doc = Document()
    style_doc(doc)

    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    bn_run(hp.add_run("মুক্তি কক্সবাজার · Odoo ১৯ — এইচআর মডিউল · প্রশিক্ষকের গাইড"), 8.5, False, GREY)

    # ---- cover -------------------------------------------------------------
    if LOGO.exists():
        lp = doc.add_paragraph()
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.add_run().add_picture(str(LOGO), width=Inches(1.35))
    para(doc, "মুক্তি কক্সবাজার", 20, True, NAVY, 2, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Mukti Cox's Bazar", 11, False, GREY, 18, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "এইচআর মডিউল — প্রশিক্ষকের গাইড", 24, True, NAVY, 4, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "HR Module — Trainer's Guide", 13, False, GREY, 16, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "একজন কর্মীর যাত্রা ধরে ধাপে ধাপে — সাতটি সেশনে",
         13, False, None, 4, WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "রাশেদুল করিমের গল্প: নিয়োগের চাহিদা থেকে চূড়ান্ত হিসাব পর্যন্ত",
         11, False, GREY, 28, WD_ALIGN_PARAGRAPH.CENTER)
    total_min = sum(s["minutes"] for s in SESSIONS)
    total_steps = sum(len(s["steps"]) for s in SESSIONS)
    table(doc, ["", ""], [
        ["সেশন সংখ্যা", f"{bn_num(len(SESSIONS))} টি"],
        ["মোট সময়", f"আনুমানিক {bn_num(total_min)} মিনিট (৩ দিনে ভাগ করা যায়)"],
        ["মোট ধাপ", f"{bn_num(total_steps)} টি স্ক্রিন"],
        ["সিস্টেমের ঠিকানা", URL],
        ["সংস্করণ", "১.০ · সেপ্টেম্বর ২০২৬"],
    ], widths=[2.0, 4.4])
    page_break(doc)

    # ---- how to use --------------------------------------------------------
    heading(doc, "কীভাবে এই গাইডটি ব্যবহার করবেন")
    para(doc, "এই গাইডটি প্রশিক্ষকের জন্য — অংশগ্রহণকারীদের হাতে দেওয়ার জন্য নয়। "
              "প্রতিটি ধাপে স্ক্রিনের ছবি আছে, আর ছবির নিচে চারটি লাইন। লাইনগুলো এভাবে পড়বেন:")
    kv(doc, "বলুন", "এটুকু আপনি মুখে বলবেন। হুবহু পড়ার দরকার নেই — ভাব ঠিক রেখে নিজের ভাষায় বলুন।", NAVY)
    kv(doc, "করুন", "স্ক্রিনে ঠিক এই জায়গায় ক্লিক করবেন। এটা ইংরেজিতে লেখা, কারণ স্ক্রিনেও ইংরেজিতেই আছে।", RED)
    kv(doc, "দেখাবে", "ক্লিকের পর যা আসবে। অংশগ্রহণকারীদের সেদিকে চোখ রাখতে বলুন।", GREEN)
    kv(doc, "খেয়াল রাখুন", "এখানেই মানুষ সবচেয়ে বেশি ভুল করে। এই লাইনটা বাদ দেবেন না।", AMBER)
    para(doc, "")
    para(doc, "একটি নিয়ম সবচেয়ে জরুরি", 12, True, NAVY, 3)
    para(doc, "বাটন বা মেনুর নাম কখনও বাংলায় অনুবাদ করবেন না। “Time Off”-কে “টাইম অফ” বলুন, "
              "“ছুটির অ্যাপ” নয়। কারণ স্ক্রিনে ইংরেজিতেই লেখা থাকবে — অনুবাদ করে বললে "
              "তাঁরা খুঁজে পাবেন না। বাকি সব কথা বাংলায় বলুন।")
    para(doc, "")
    para(doc, "দুটি ধারা — কে কোন সেশনে থাকবেন", 12, True, NAVY, 3)
    table(doc, ["ধারা", "কারা", "কোন সেশন"], [
        ["ধারা ১ — একবার সেট-আপ", "এইচআর, ফিন্যান্স, ব্যবস্থাপক, প্রধান নির্বাহী",
         "সেশন ১, ২, ৩, ৬, ৭"],
        ["ধারা ২ — প্রতিদিনের কাজ", "প্রতিষ্ঠানের সব কর্মী", "সেশন ৪ ও ৫ (বাধ্যতামূলক)"],
    ], widths=[1.8, 2.6, 2.0])
    para(doc, "")
    para(doc, "প্রতিটি সেশনের শুরুতে বলে দিন এই সেশনটা কাদের জন্য। যাঁদের দরকার নেই তাঁরা "
              "বসে থাকলে মনোযোগ হারান, আর যাঁদের দরকার তাঁরা পিছিয়ে পড়েন।", 10.5, False, GREY)
    page_break(doc)

    # ---- before you train ---------------------------------------------------
    heading(doc, "প্রশিক্ষণের আগে")
    para(doc, "ক. প্রস্তুতির তালিকা", 12, True, NAVY, 4)
    for i, item in enumerate([
        "প্রজেক্টর বা বড় স্ক্রিন ঠিকমতো কাজ করছে কি না দেখে নিন — লেখা যেন পেছনের সারি থেকেও পড়া যায়।",
        f"ইন্টারনেট পরীক্ষা করুন এবং ঠিকানাটা আগে থেকে ব্রাউজারে খুলে রাখুন: {URL}",
        "প্রশিক্ষণের আগের দিন নিজে একবার পুরো পথটা হেঁটে দেখুন — কোথাও আটকে গেলে আগেই জানা থাকবে।",
        "প্রত্যেক অংশগ্রহণকারীর জন্য একটি করে লগইন আগে থেকে তৈরি রাখুন। একজনের লগইন দিয়ে অনেকে কাজ করলে শেখা হয় না।",
        "সেশন ৪ ও ৫-এ দুইজনের লগইন লাগবে — একজন কর্মী, একজন অনুমোদনকারী। দুটোই হাতের কাছে রাখুন।",
        "হাতে-কলমে অনুশীলনের জন্য অন্তত ৩০ মিনিট আলাদা রাখুন। শুধু দেখিয়ে গেলে কেউ মনে রাখে না।",
    ], 1):
        p = doc.add_paragraph(style="List Number")
        bn_run(p.add_run(item), 10.5)
    para(doc, "")

    para(doc, "খ. যা আগে ঠিক করে নিতে হবে", 12, True, RED, 4)
    para(doc, "নিচের বিষয়গুলো সিস্টেমে এখনো অমীমাংসিত। প্রশিক্ষণে এগুলো সামনে চলে এলে "
              "যেন অপ্রস্তুত না হন, তাই আগেই জানিয়ে রাখা হলো।", 10.5, False, GREY)
    table(doc, ["বিষয়", "এখন যা অবস্থা", "প্রশিক্ষকের করণীয়"],
          BEFORE_YOU_TRAIN, widths=[1.5, 2.6, 2.3], sizes=(9, 9))
    page_break(doc)

    # ---- cast + logins ------------------------------------------------------
    heading(doc, "গল্পের চরিত্র ও লগইন")
    para(doc, "পুরো প্রশিক্ষণে আমরা একজন মানুষকেই অনুসরণ করব। একেকটা সেশনে একেক মেনু "
              "ঘুরে দেখানোর চেয়ে একজনের যাত্রা ধরে এগোলে মানুষ অনেক বেশি মনে রাখে।")
    table(doc, ["চরিত্র", "পরিচয়", "গল্পে ভূমিকা"], [
        ["রাশেদুল করিম", "MEAL Officer · কক্সবাজার · কর্মী কোড MEAL-001 · গ্রেড G6",
         "যাঁর যাত্রা আমরা শুরু থেকে শেষ পর্যন্ত দেখব"],
        ["ফারজানা আক্তার", "কর্মী কোড CORE-003", "যিনি নিয়মের চেয়ে কম নোটিশে চাকরি ছাড়েন (সেশন ৭)"],
        ["প্রকল্প ব্যবস্থাপক", "MEAL টিমের দায়িত্বে", "পদের চাহিদা তোলেন, ছুটি ও ভ্রমণ অনুমোদন করেন"],
        ["এইচআর", "মানবসম্পদ বিভাগ", "যাচাই করে, ফাইল খোলে, যোগদান ও বিদায় সামলায়"],
        ["প্রধান নির্বাহী (CE)", "সর্বোচ্চ অনুমোদনকারী", "চাহিদা ও বড় সিদ্ধান্তে চূড়ান্ত অনুমোদন"],
        ["ফিন্যান্স", "হিসাব বিভাগ", "বেতন করে, ভ্রমণ ভাতার টাকা ছাড়ে"],
    ], widths=[1.35, 2.6, 2.45], sizes=(9.5, 9))
    para(doc, "")
    para(doc, "প্রশিক্ষণের লগইন", 12, True, NAVY, 4)
    table(doc, ["কার ভূমিকায়", "Username", "Password"], [
        ["কর্মী — রাশেদুল করিম", "rashedul", "rashedul"],
        ["এইচআর / ম্যানেজার / প্রধান নির্বাহী", "admin", "admin"],
    ], widths=[3.0, 1.7, 1.7])
    para(doc, "")
    para(doc, "একই ব্রাউজারে দুইজনের লগইন একসাথে চলে না। কর্মী থেকে অনুমোদনকারীতে যেতে হলে "
              "লগআউট করুন, অথবা আরেকটি ব্রাউজার উইন্ডো (Incognito / Private) খুলে নিন — "
              "প্রশিক্ষণে এটাই সবচেয়ে সহজ।", 10.5, False, AMBER)
    page_break(doc)

    # ---- agenda -------------------------------------------------------------
    heading(doc, "প্রশিক্ষণ সূচি")
    rows = []
    for s in SESSIONS:
        rows.append([f"সেশন {bn_num(s['no'])}", s["bn"], s["who"], f"{bn_num(s['minutes'])} মিনিট"])
    table(doc, ["", "বিষয়", "কারা থাকবেন", "সময়"], rows,
          widths=[0.85, 2.2, 2.4, 0.95], sizes=(9.5, 9.5))
    para(doc, "")
    para(doc, "তিন দিনে ভাগ করার প্রস্তাব", 12, True, NAVY, 4)
    table(doc, ["দিন", "সেশন", "কারা"], [
        ["দিন ১", "সেশন ১–৩ — নিয়োগ থেকে যোগদান", "এইচআর, ব্যবস্থাপক, প্রধান নির্বাহী"],
        ["দিন ২", "সেশন ৪–৫ — ছুটি ও ভ্রমণ ভাতা", "প্রতিষ্ঠানের সব কর্মী (বাধ্যতামূলক)"],
        ["দিন ৩", "সেশন ৬–৭ — বেতন ও বিদায়", "ফিন্যান্স ও এইচআর"],
    ], widths=[0.8, 3.0, 2.6])
    para(doc, "")
    para(doc, "প্রতিটি সেশনের জন্য একটি করে ভিডিও আছে। চাইলে সেশন শুরুর আগে ভিডিওটা একবার "
              "চালিয়ে দিতে পারেন, তারপর লাইভ দেখাতে পারেন — অথবা উল্টোটা। ভিডিওগুলো আছে "
              "mukticox/video ফোল্ডারে।", 10.5, False, GREY)
    page_break(doc)

    # ---- policy at a glance -------------------------------------------------
    heading(doc, "MCB নীতিমালা — এক নজরে")
    para(doc, "প্রশিক্ষণে সংখ্যা নিয়ে প্রশ্ন আসবেই। নিচের তথ্যগুলো সিস্টেমের কোড থেকে "
              "যাচাই করা — এগুলো নির্ভয়ে বলতে পারেন। যেগুলো এখনো নিশ্চিত নয়, সেগুলো "
              "আলাদা করে চিহ্নিত করা আছে।", 10.5, False, GREY)
    for title, headers, rows, widths in POLICY:
        para(doc, "")
        para(doc, title, 12, True, NAVY, 4)
        table(doc, headers, rows, widths=widths, sizes=(9, 9))
    page_break(doc)

    # ---- the sessions -------------------------------------------------------
    for s in SESSIONS:
        heading(doc, f"সেশন {bn_num(s['no'])} — {s['bn']}")
        para(doc, s["en"], 11, False, GREY, 8)

        table(doc, ["কারা থাকবেন", "সময়", "ভিডিও"],
              [[s["who"], f"{bn_num(s['minutes'])} মিনিট", s["video"]]],
              widths=[2.8, 1.2, 2.4], sizes=(9, 9))
        para(doc, "")

        para(doc, "গল্পটা", 12, True, NAVY, 3)
        para(doc, s["story"], 11)
        para(doc, "")

        para(doc, "এই সেশনে যা শেখানো হবে", 12, True, NAVY, 3)
        for g in s["goals"]:
            p = doc.add_paragraph(style="List Bullet")
            bn_run(p.add_run(g), 10.5)
        para(doc, "")

        # steps
        steps_file = FLOWS / s["key"] / "steps.json"
        steps = json.loads(steps_file.read_text()) if steps_file.exists() else []
        narr = s["steps"]
        used = set()
        para(doc, "ধাপে ধাপে", 12, True, NAVY, 6)
        for st in steps:
            slug = st["slug"]
            # a flow can log in twice (employee, then approver) — the second login
            # screen reuses the slug, so key the narration as slug_2 for that one
            key = slug if slug not in used else f"{slug}_2"
            used.add(slug)
            n = narr.get(key) or narr.get(slug)
            if not n:
                continue
            hp2 = doc.add_paragraph()
            hp2.paragraph_format.space_before = Pt(10)
            hp2.paragraph_format.space_after = Pt(4)
            hp2.paragraph_format.keep_with_next = True
            hp2.paragraph_format.page_break_before = False
            bn_run(hp2.add_run(f"ধাপ {bn_num(st['n'])}  ·  "), 11, True, NAVY)
            bn_run(hp2.add_run(st["headline"]), 11, True, GREY)
            embed(doc, FLOWS / s["key"] / st["file"])
            kv(doc, "বলুন", n["say"], NAVY)
            kv(doc, "করুন", n["do"], RED)
            kv(doc, "দেখাবে", n["see"], GREEN)
            if n.get("tip"):
                kv(doc, "খেয়াল রাখুন", n["tip"], AMBER)

        # wrap-up
        para(doc, "")
        para(doc, "যেখানে সবচেয়ে বেশি ভুল হয়", 12, True, RED, 3)
        for m in s["mistakes"]:
            p = doc.add_paragraph(style="List Bullet")
            bn_run(p.add_run(m), 10.5)
        para(doc, "")
        para(doc, "যাচাই প্রশ্ন — সেশন শেষে জিজ্ঞেস করুন", 12, True, NAVY, 3)
        for q in s["questions"]:
            p = doc.add_paragraph(style="List Number")
            bn_run(p.add_run(q), 10.5)
        para(doc, "")
        para(doc, "হাতে-কলমে অনুশীলন", 12, True, GREEN, 3)
        para(doc, s["practice"], 10.5)
        page_break(doc)

    # ---- quick reference ----------------------------------------------------
    heading(doc, "দ্রুত রেফারেন্স — কর্মীদের হাতে দেওয়ার জন্য")
    para(doc, "এই পাতাটা আলাদা করে প্রিন্ট করে সবাইকে দিয়ে দিতে পারেন।", 10.5, False, GREY)
    para(doc, "")
    table(doc, ["যা করতে চান", "কোথায় যাবেন"], QUICKREF, widths=[2.7, 3.7])
    para(doc, "")
    para(doc, "মনে রাখার তিনটি কথা", 12, True, NAVY, 4)
    for t in [
        "ছুটির আবেদনের আগে সবসময় ব্যালেন্স দেখে নিন।",
        "ভ্রমণ ভাতার টাকা নিজে হাতে লিখবেন না — Apply Per Diem চাপুন, হার নিজে বসবে।",
        "কাজ শেষে সবসময় লগআউট করুন। আপনার লগইনে যা হবে, তার দায় আপনার।",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        bn_run(p.add_run(t), 11)
    page_break(doc)

    # ---- faq ----------------------------------------------------------------
    heading(doc, "সাধারণ প্রশ্ন ও উত্তর")
    para(doc, "প্রশিক্ষণে এই প্রশ্নগুলো প্রায় সবসময়ই আসে।", 10.5, False, GREY)
    for q, a in FAQ:
        para(doc, "")
        para(doc, f"প্রশ্ন: {q}", 11, True, NAVY, 2)
        para(doc, f"উত্তর: {a}", 10.5)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size:,} bytes)")
    imgs = sum(len(json.loads((FLOWS / s['key'] / 'steps.json').read_text()))
               for s in SESSIONS if (FLOWS / s['key'] / 'steps.json').exists())
    print(f"  sessions: {len(SESSIONS)} · step screens available: {imgs}")
    print(f"  missing images: {MISSING or 'none'}")


if __name__ == "__main__":
    build()
